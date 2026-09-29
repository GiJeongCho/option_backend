from __future__ import annotations

import bisect
import csv
import json
import math
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

import minimal_baseline_backtest as base


OUTPUT_DIR = Path(__file__).resolve().parent / "output"
VOLATILITY_PATH = (
    Path(__file__).resolve().parents[2]
    / "rowdata"
    / "코스피200 변동성지수.xlsx"
)

STRATEGY_VERSION = "S1-v0.1"
INITIAL_CASH = 80_000_000
DAILY_PRINCIPAL_LIMIT = 5_000_000
MULTIPLIER = 250_000
COMMISSION = 500
SLIPPAGE = 0.01
VOLUME_PARTICIPATION = 0.10
ENTRY_WAIT_MINUTES = 5
INITIAL_ENTRY_CUTOFF = 1400
FORCE_EXIT_TIME = 1515
PREMIUM_MIN = 1.00
PREMIUM_MAX = 1.49
VIX_THRESHOLD = 30.0
TRANCHE_BUDGETS = {1: 1_250_000, 2: 1_750_000, 3: 2_000_000}
REENTRY_COOLDOWN_MINUTES = 20
MAX_CAMPAIGNS = 2


@dataclass(frozen=True)
class EntrySignal:
    code: str
    index: int
    minute: int
    close: float
    relative_strength: float
    recent_volume: int
    setup_low: float


@dataclass
class Campaign:
    campaign_id: int
    code: str
    name: str
    call_put: str
    setup_low: float
    first_signal_minute: int
    first_entry_minute: int
    quantity: int = 0
    total_bought_quantity: int = 0
    weighted_entry_value: float = 0.0
    buy_principal: int = 0
    buy_fee: int = 0
    gross_sales: int = 0
    sell_fee: int = 0
    cash_flow: int = 0
    stage: int = 0
    pullback_seen: bool = False
    first_target_done: bool = False
    second_target_done: bool = False
    high_water: float = 0.0
    below_ma_count: int = 0
    mfe_pct: float = 0.0
    mae_pct: float = 0.0
    exit_reason: str = ""
    exit_minute: int | None = None
    expired_quantity: int = 0
    trades: list[dict[str, object]] = field(default_factory=list)

    @property
    def average_entry(self) -> float:
        if self.total_bought_quantity <= 0:
            return 0.0
        return self.weighted_entry_value / self.total_bought_quantity


@dataclass
class DayResult:
    date: str
    previous_vix: float
    high_volatility: bool
    start_cash: int
    end_cash: int
    pnl: int
    signal_found: bool
    trade_day: bool
    campaign_count: int
    profitable_campaigns: int
    losing_campaigns: int
    buy_principal: int
    fees: int
    gross_sales: int
    expired_quantity: int
    first_entry_minute: int | None
    call_put: str
    exit_reasons: str
    mfe_pct: float
    mae_pct: float


def volatility_closes() -> tuple[list[str], list[float]]:
    workbook = load_workbook(VOLATILITY_PATH, read_only=True, data_only=True)
    try:
        sheet = workbook["코스피200 변동성지수"]
        rows: list[tuple[str, float]] = []
        for raw_date, _, _, _, raw_close in sheet.iter_rows(
            min_row=2, max_col=5, values_only=True
        ):
            if raw_date is None or raw_close is None:
                continue
            if isinstance(raw_date, datetime):
                value = raw_date.date()
            elif isinstance(raw_date, date):
                value = raw_date
            else:
                text = str(raw_date).strip()
                value = None
                for date_format in ("%Y-%m-%d", "%Y%m%d", "%Y.%m.%d"):
                    try:
                        value = datetime.strptime(text, date_format).date()
                        break
                    except ValueError:
                        continue
                if value is None:
                    continue
            rows.append((value.strftime("%Y%m%d"), float(raw_close)))
        rows.sort()
        return [item[0] for item in rows], [item[1] for item in rows]
    finally:
        workbook.close()


def previous_volatility(
    trading_date: str, dates: list[str], closes: list[float]
) -> float | None:
    index = bisect.bisect_left(dates, trading_date) - 1
    return closes[index] if index >= 0 else None


def prior_average_volume(contract: base.ContractGrid, index: int) -> float:
    return (
        sum(
            contract.bars[position].volume
            for position in range(index - 20, index)
        )
        / 20
    )


def prior_high(
    contract: base.ContractGrid, index: int, period: int
) -> float | None:
    values = [
        contract.bars[position].high
        for position in range(index - period, index)
    ]
    if any(value is None for value in values):
        return None
    return max(float(value) for value in values)


def prior_low(
    contract: base.ContractGrid, index: int, period: int
) -> float | None:
    values = [
        contract.bars[position].low
        for position in range(index - period, index)
    ]
    if any(value is None for value in values):
        return None
    return min(float(value) for value in values)


def initial_candidates(
    contracts: dict[str, base.ContractGrid], index: int
) -> list[EntrySignal]:
    if index < 20:
        return []
    candidates: list[EntrySignal] = []
    for contract in contracts.values():
        bar = contract.bars[index]
        previous = contract.bars[index - 1]
        ma5 = contract.ma5[index]
        previous_ma5 = contract.ma5[index - 1]
        ma20 = contract.ma20[index]
        setup_low = prior_low(contract, index, 10)
        if (
            not bar.traded
            or bar.close is None
            or previous.close is None
            or ma5 is None
            or previous_ma5 is None
            or ma20 is None
            or setup_low is None
            or not PREMIUM_MIN <= bar.close <= PREMIUM_MAX
        ):
            continue
        average_volume = prior_average_volume(contract, index)
        if not (
            previous.close <= previous_ma5
            and bar.close > ma5
            and bar.volume > average_volume
        ):
            continue
        candidates.append(
            EntrySignal(
                code=contract.code,
                index=index,
                minute=bar.minute,
                close=bar.close,
                relative_strength=(bar.close / ma20) - 1.0,
                recent_volume=sum(
                    item.volume
                    for item in contract.bars[index - 19 : index + 1]
                ),
                setup_low=setup_low,
            )
        )
    return sorted(
        candidates,
        key=lambda item: (
            -item.relative_strength,
            -item.recent_volume,
            item.code,
        ),
    )


def next_traded_bar(
    contract: base.ContractGrid, signal_index: int
) -> tuple[int, base.MinuteBar] | None:
    last_index = min(
        len(contract.bars) - 1, signal_index + ENTRY_WAIT_MINUTES
    )
    for index in range(signal_index + 1, last_index + 1):
        bar = contract.bars[index]
        if bar.traded:
            return index, bar
    return None


def buy_tranche(
    date_value: str,
    campaign: Campaign,
    contract: base.ContractGrid,
    signal_index: int,
    signal_close: float,
    stage: int,
    budget: int,
    available_cash: int,
    remaining_daily_principal: int,
) -> tuple[int, int] | None:
    fill = next_traded_bar(contract, signal_index)
    if fill is None:
        return None
    fill_index, bar = fill
    assert bar.open is not None
    fill_price = bar.open + SLIPPAGE
    allowed_budget = min(budget, remaining_daily_principal)
    requested = math.floor(
        allowed_budget / ((signal_close + SLIPPAGE) * MULTIPLIER)
    )
    by_actual = math.floor(allowed_budget / (fill_price * MULTIPLIER))
    by_cash = math.floor(
        available_cash
        / (fill_price * MULTIPLIER + 2 * COMMISSION)
    )
    by_volume = math.floor(bar.volume * VOLUME_PARTICIPATION)
    quantity = min(requested, by_actual, by_cash, by_volume)
    if quantity <= 0:
        return None

    principal = round(quantity * fill_price * MULTIPLIER)
    fee = quantity * COMMISSION
    campaign.quantity += quantity
    campaign.total_bought_quantity += quantity
    campaign.weighted_entry_value += fill_price * quantity
    campaign.buy_principal += principal
    campaign.buy_fee += fee
    campaign.cash_flow -= principal + fee
    campaign.stage = max(campaign.stage, stage)
    campaign.high_water = max(campaign.high_water, fill_price)
    campaign.trades.append(
        {
            "strategy": STRATEGY_VERSION,
            "date": date_value,
            "campaign": campaign.campaign_id,
            "minute": bar.minute,
            "code": contract.code,
            "call_put": contract.call_put,
            "side": "BUY",
            "reason": f"ENTRY_STAGE_{stage}",
            "stage": stage,
            "quantity": quantity,
            "price": fill_price,
            "principal": principal,
            "fee": fee,
            "cash_flow": -(principal + fee),
            "bar_volume": bar.volume,
        }
    )
    return fill_index, principal


def sell_order(
    date_value: str,
    campaign: Campaign,
    contract: base.ContractGrid,
    signal_index: int,
    requested_quantity: int,
    reason: str,
) -> tuple[int, int, bool]:
    remaining = min(requested_quantity, campaign.quantity)
    requested = remaining
    capacity = 0.0
    last_fill_index = signal_index
    for index in range(signal_index + 1, len(contract.bars)):
        bar = contract.bars[index]
        if not bar.traded:
            continue
        assert bar.open is not None
        capacity += bar.volume * VOLUME_PARTICIPATION
        fillable = min(remaining, math.floor(capacity + 1e-12))
        if fillable <= 0:
            continue
        capacity -= fillable
        fill_price = max(0.0, bar.open - SLIPPAGE)
        principal = round(fillable * fill_price * MULTIPLIER)
        fee = fillable * COMMISSION
        campaign.quantity -= fillable
        campaign.gross_sales += principal
        campaign.sell_fee += fee
        campaign.cash_flow += principal - fee
        remaining -= fillable
        last_fill_index = index
        campaign.trades.append(
            {
                "strategy": STRATEGY_VERSION,
                "date": date_value,
                "campaign": campaign.campaign_id,
                "minute": bar.minute,
                "code": contract.code,
                "call_put": contract.call_put,
                "side": "SELL",
                "reason": reason,
                "stage": campaign.stage,
                "quantity": fillable,
                "price": fill_price,
                "principal": principal,
                "fee": fee,
                "cash_flow": principal - fee,
                "bar_volume": bar.volume,
            }
        )
        if remaining <= 0:
            break
    return requested - remaining, last_fill_index, remaining <= 0


def expire_remaining(
    date_value: str, campaign: Campaign, reason: str
) -> None:
    if campaign.quantity <= 0:
        return
    quantity = campaign.quantity
    campaign.quantity = 0
    campaign.expired_quantity += quantity
    campaign.trades.append(
        {
            "strategy": STRATEGY_VERSION,
            "date": date_value,
            "campaign": campaign.campaign_id,
            "minute": base.SESSION_MINUTES[-1],
            "code": campaign.code,
            "call_put": campaign.call_put,
            "side": "EXPIRE",
            "reason": reason,
            "stage": campaign.stage,
            "quantity": quantity,
            "price": 0.0,
            "principal": 0,
            "fee": 0,
            "cash_flow": 0,
            "bar_volume": 0,
        }
    )


def update_excursions(campaign: Campaign, close: float) -> None:
    average = campaign.average_entry
    if average <= 0:
        return
    change = close / average - 1.0
    campaign.mfe_pct = max(campaign.mfe_pct, change * 100)
    campaign.mae_pct = min(campaign.mae_pct, change * 100)
    campaign.high_water = max(campaign.high_water, close)


def close_campaign(
    campaign: Campaign, reason: str, minute: int
) -> None:
    campaign.exit_reason = reason
    campaign.exit_minute = minute


def campaign_row(date_value: str, campaign: Campaign) -> dict[str, object]:
    return {
        "date": date_value,
        "campaign": campaign.campaign_id,
        "code": campaign.code,
        "name": campaign.name,
        "call_put": campaign.call_put,
        "first_signal_minute": campaign.first_signal_minute,
        "first_entry_minute": campaign.first_entry_minute,
        "exit_minute": campaign.exit_minute,
        "tranche_count": campaign.stage,
        "total_bought_quantity": campaign.total_bought_quantity,
        "buy_principal": campaign.buy_principal,
        "buy_fee": campaign.buy_fee,
        "gross_sales": campaign.gross_sales,
        "sell_fee": campaign.sell_fee,
        "pnl": campaign.cash_flow,
        "mfe_pct": round(campaign.mfe_pct, 4),
        "mae_pct": round(campaign.mae_pct, 4),
        "exit_reason": campaign.exit_reason,
        "expired_quantity": campaign.expired_quantity,
        "setup_low": campaign.setup_low,
        "average_entry": round(campaign.average_entry, 6),
    }


def simulate_day(
    date_value: str,
    contracts: dict[str, base.ContractGrid],
    previous_vix: float,
    start_cash: int,
) -> tuple[DayResult, list[dict[str, object]], list[dict[str, object]]]:
    high_volatility = previous_vix >= VIX_THRESHOLD
    if not high_volatility:
        return (
            DayResult(
                date_value,
                previous_vix,
                False,
                start_cash,
                start_cash,
                0,
                False,
                False,
                0,
                0,
                0,
                0,
                0,
                0,
                0,
                None,
                "",
                "",
                0.0,
                0.0,
            ),
            [],
            [],
        )

    active: Campaign | None = None
    completed: list[Campaign] = []
    all_trades: list[dict[str, object]] = []
    daily_principal = 0
    day_cash_flow = 0
    signal_found = False
    next_initial_index = 20
    index = 20

    while index < len(base.SESSION_MINUTES):
        if active is None:
            if (
                len(completed) >= MAX_CAMPAIGNS
                or daily_principal >= DAILY_PRINCIPAL_LIMIT
                or index < next_initial_index
                or base.SESSION_MINUTES[index] > INITIAL_ENTRY_CUTOFF
            ):
                index += 1
                continue
            candidates = initial_candidates(contracts, index)
            if not candidates:
                index += 1
                continue
            signal_found = True
            signal = candidates[0]
            contract = contracts[signal.code]
            campaign = Campaign(
                campaign_id=len(completed) + 1,
                code=contract.code,
                name=contract.name,
                call_put=contract.call_put,
                setup_low=signal.setup_low,
                first_signal_minute=signal.minute,
                first_entry_minute=0,
            )
            result = buy_tranche(
                date_value,
                campaign,
                contract,
                signal.index,
                signal.close,
                1,
                TRANCHE_BUDGETS[1],
                start_cash + day_cash_flow,
                DAILY_PRINCIPAL_LIMIT - daily_principal,
            )
            if result is None:
                index += ENTRY_WAIT_MINUTES
                continue
            fill_index, principal = result
            campaign.first_entry_minute = contract.bars[fill_index].minute
            daily_principal += principal
            day_cash_flow += campaign.trades[-1]["cash_flow"]  # type: ignore[operator]
            active = campaign
            index = fill_index
            continue

        contract = contracts[active.code]
        bar = contract.bars[index]
        ma5 = contract.ma5[index]
        ma20 = contract.ma20[index]
        if bar.traded and bar.close is not None:
            update_excursions(active, bar.close)
            if ma5 is not None and ma20 is not None and ma5 < ma20:
                active.below_ma_count += 1
            else:
                active.below_ma_count = 0

            average = active.average_entry
            stop_reason = ""
            if bar.close <= average * 0.70:
                stop_reason = "STOP_30_PERCENT"
            elif bar.close < active.setup_low:
                stop_reason = "STOP_SETUP_LOW"
            elif active.below_ma_count >= 2:
                stop_reason = "STOP_MA_REVERSAL"

            if stop_reason:
                filled, fill_index, complete = sell_order(
                    date_value,
                    active,
                    contract,
                    index,
                    active.quantity,
                    stop_reason,
                )
                if filled:
                    day_cash_flow = (
                        sum(item.cash_flow for item in completed)
                        + active.cash_flow
                    )
                if not complete:
                    expire_remaining(
                        date_value, active, "UNFILLED_STOP_ZERO_VALUE"
                    )
                close_campaign(
                    active,
                    stop_reason if complete else "UNFILLED_STOP_ZERO_VALUE",
                    (
                        contract.bars[fill_index].minute
                        if complete
                        else base.SESSION_MINUTES[-1]
                    ),
                )
                all_trades.extend(active.trades)
                completed.append(active)
                active = None
                next_initial_index = fill_index + REENTRY_COOLDOWN_MINUTES
                index = fill_index + 1
                continue

            if (
                not active.first_target_done
                and bar.close >= average * 2.0
            ):
                net_per_contract = (
                    bar.close * MULTIPLIER - COMMISSION
                )
                recover = math.ceil(
                    (active.buy_principal + active.buy_fee)
                    / max(net_per_contract, 1)
                )
                request = min(active.quantity, max(1, recover))
                filled, fill_index, complete = sell_order(
                    date_value,
                    active,
                    contract,
                    index,
                    request,
                    "TARGET_2X_RECOVER",
                )
                active.first_target_done = complete
                day_cash_flow = sum(item.cash_flow for item in completed) + active.cash_flow
                if not complete:
                    expire_remaining(
                        date_value, active, "UNFILLED_TARGET_ZERO_VALUE"
                    )
                    close_campaign(
                        active,
                        "UNFILLED_TARGET_ZERO_VALUE",
                        base.SESSION_MINUTES[-1],
                    )
                    all_trades.extend(active.trades)
                    completed.append(active)
                    active = None
                    index = len(base.SESSION_MINUTES)
                    continue
                if active.quantity <= 0:
                    close_campaign(
                        active,
                        "TARGET_2X_RECOVER",
                        contract.bars[fill_index].minute,
                    )
                    all_trades.extend(active.trades)
                    completed.append(active)
                    active = None
                    next_initial_index = (
                        fill_index + REENTRY_COOLDOWN_MINUTES
                    )
                    index = fill_index + 1
                    continue
                index = fill_index
                continue

            if (
                active.first_target_done
                and not active.second_target_done
                and bar.close >= average * 5.0
            ):
                active.second_target_done = True
                request = (
                    min(
                        active.quantity - 1,
                        max(1, math.ceil(active.quantity / 2)),
                    )
                    if active.quantity > 1
                    else 0
                )
                if request > 0:
                    _, fill_index, complete = sell_order(
                        date_value,
                        active,
                        contract,
                        index,
                        request,
                        "TARGET_5X_PARTIAL",
                    )
                    day_cash_flow = (
                        sum(item.cash_flow for item in completed)
                        + active.cash_flow
                    )
                    if not complete:
                        expire_remaining(
                            date_value,
                            active,
                            "UNFILLED_TARGET_ZERO_VALUE",
                        )
                        close_campaign(
                            active,
                            "UNFILLED_TARGET_ZERO_VALUE",
                            base.SESSION_MINUTES[-1],
                        )
                        all_trades.extend(active.trades)
                        completed.append(active)
                        active = None
                        index = len(base.SESSION_MINUTES)
                        continue
                    index = fill_index
                    continue

            if active.first_target_done and (
                bar.close <= active.high_water * 0.70
                or (ma20 is not None and bar.close < ma20)
            ):
                _, fill_index, complete = sell_order(
                    date_value,
                    active,
                    contract,
                    index,
                    active.quantity,
                    "TRAIL_OR_MA20",
                )
                day_cash_flow = sum(item.cash_flow for item in completed) + active.cash_flow
                if not complete:
                    expire_remaining(
                        date_value, active, "UNFILLED_TRAIL_ZERO_VALUE"
                    )
                close_campaign(
                    active,
                    "TRAIL_OR_MA20"
                    if complete
                    else "UNFILLED_TRAIL_ZERO_VALUE",
                    (
                        contract.bars[fill_index].minute
                        if complete
                        else base.SESSION_MINUTES[-1]
                    ),
                )
                all_trades.extend(active.trades)
                completed.append(active)
                active = None
                next_initial_index = fill_index + REENTRY_COOLDOWN_MINUTES
                index = fill_index + 1
                continue

            if active.stage == 1:
                previous_high = prior_high(contract, index, 5)
                if (
                    bar.minute < FORCE_EXIT_TIME
                    and ma5 is not None
                    and ma20 is not None
                    and previous_high is not None
                    and ma5 > ma20
                    and bar.close > previous_high
                    and bar.close > average
                    and bar.volume > prior_average_volume(contract, index)
                ):
                    before = len(active.trades)
                    result = buy_tranche(
                        date_value,
                        active,
                        contract,
                        index,
                        bar.close,
                        2,
                        TRANCHE_BUDGETS[2],
                        start_cash + day_cash_flow,
                        DAILY_PRINCIPAL_LIMIT - daily_principal,
                    )
                    if result is not None:
                        fill_index, principal = result
                        daily_principal += principal
                        day_cash_flow += sum(
                            int(item["cash_flow"])
                            for item in active.trades[before:]
                        )
                        index = fill_index
                        continue

            if (
                active.stage == 2
                and active.pullback_seen
                and bar.minute < FORCE_EXIT_TIME
                and ma5 is not None
                and index > 0
                and contract.bars[index - 1].high is not None
                and bar.close > ma5
                and bar.close > float(contract.bars[index - 1].high)
                and bar.close > average
            ):
                before = len(active.trades)
                result = buy_tranche(
                    date_value,
                    active,
                    contract,
                    index,
                    bar.close,
                    3,
                    TRANCHE_BUDGETS[3],
                    start_cash + day_cash_flow,
                    DAILY_PRINCIPAL_LIMIT - daily_principal,
                )
                if result is not None:
                    fill_index, principal = result
                    daily_principal += principal
                    day_cash_flow += sum(
                        int(item["cash_flow"])
                        for item in active.trades[before:]
                    )
                    index = fill_index
                    continue

            if (
                active.stage >= 2
                and ma5 is not None
                and ma20 is not None
                and bar.low is not None
                and bar.low <= ma5 * 1.01
                and bar.close >= ma20
            ):
                active.pullback_seen = True

        if bar.minute >= FORCE_EXIT_TIME:
            _, fill_index, complete = sell_order(
                date_value,
                active,
                contract,
                index,
                active.quantity,
                "TIME",
            )
            day_cash_flow = sum(item.cash_flow for item in completed) + active.cash_flow
            if not complete:
                expire_remaining(
                    date_value, active, "UNFILLED_TIME_ZERO_VALUE"
                )
            close_campaign(
                active,
                "TIME" if complete else "UNFILLED_TIME_ZERO_VALUE",
                (
                    contract.bars[fill_index].minute
                    if complete
                    else base.SESSION_MINUTES[-1]
                ),
            )
            all_trades.extend(active.trades)
            completed.append(active)
            active = None
            break
        index += 1

    if active is not None:
        expire_remaining(date_value, active, "DATA_END_ZERO_VALUE")
        close_campaign(active, "DATA_END_ZERO_VALUE", base.SESSION_MINUTES[-1])
        all_trades.extend(active.trades)
        completed.append(active)

    day_cash_flow = sum(item.cash_flow for item in completed)
    campaign_rows = [campaign_row(date_value, item) for item in completed]
    fees = sum(item.buy_fee + item.sell_fee for item in completed)
    call_put_values = {item.call_put for item in completed}
    day = DayResult(
        date=date_value,
        previous_vix=previous_vix,
        high_volatility=True,
        start_cash=start_cash,
        end_cash=start_cash + day_cash_flow,
        pnl=day_cash_flow,
        signal_found=signal_found,
        trade_day=bool(completed),
        campaign_count=len(completed),
        profitable_campaigns=sum(item.cash_flow > 0 for item in completed),
        losing_campaigns=sum(item.cash_flow < 0 for item in completed),
        buy_principal=sum(item.buy_principal for item in completed),
        fees=fees,
        gross_sales=sum(item.gross_sales for item in completed),
        expired_quantity=sum(item.expired_quantity for item in completed),
        first_entry_minute=(
            min(item.first_entry_minute for item in completed)
            if completed
            else None
        ),
        call_put=(
            next(iter(call_put_values))
            if len(call_put_values) == 1
            else "MIXED"
            if call_put_values
            else ""
        ),
        exit_reasons=",".join(item.exit_reason for item in completed),
        mfe_pct=max((item.mfe_pct for item in completed), default=0.0),
        mae_pct=min((item.mae_pct for item in completed), default=0.0),
    )
    return day, campaign_rows, all_trades


def max_drawdown(values: list[int]) -> int:
    peak = values[0]
    worst = 0
    for value in values:
        peak = max(peak, value)
        worst = min(worst, value - peak)
    return worst


def grouped_losses(
    rows: list[dict[str, object]], key_name: str
) -> list[dict[str, object]]:
    grouped: dict[str, list[int]] = defaultdict(list)
    for row in rows:
        if int(row["pnl"]) >= 0:
            continue
        grouped[str(row[key_name])].append(int(row["pnl"]))
    return sorted(
        (
            {
                "group": key,
                "loss_count": len(values),
                "total_loss": sum(values),
                "average_loss": round(sum(values) / len(values), 2),
                "worst_loss": min(values),
            }
            for key, values in grouped.items()
        ),
        key=lambda item: int(item["total_loss"]),
    )


def entry_time_bucket(value: int) -> str:
    if value < 1000:
        return "08:45-09:59"
    if value < 1100:
        return "10:00-10:59"
    if value < 1200:
        return "11:00-11:59"
    if value < 1300:
        return "12:00-12:59"
    return "13:00-14:00"


def vix_bucket(value: float) -> str:
    if value < 40:
        return "30-39.99"
    if value < 50:
        return "40-49.99"
    return "50+"


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8-sig")
        return
    with path.open("w", encoding="utf-8-sig", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def audit(
    daily: list[dict[str, object]],
    campaigns: list[dict[str, object]],
    trades: list[dict[str, object]],
) -> dict[str, object]:
    errors: list[str] = []
    previous_cash = INITIAL_CASH
    trades_by_day: dict[str, list[dict[str, object]]] = defaultdict(list)
    for trade in trades:
        trades_by_day[str(trade["date"])].append(trade)
    campaigns_by_day: dict[str, list[dict[str, object]]] = defaultdict(list)
    for campaign in campaigns:
        campaigns_by_day[str(campaign["date"])].append(campaign)

    for row in daily:
        date_value = str(row["date"])
        if int(row["start_cash"]) != previous_cash:
            errors.append(f"{date_value}: cash chain")
        if int(row["end_cash"]) != int(row["start_cash"]) + int(row["pnl"]):
            errors.append(f"{date_value}: pnl equation")
        if int(row["buy_principal"]) > DAILY_PRINCIPAL_LIMIT:
            errors.append(f"{date_value}: principal limit")
        if sum(int(item["cash_flow"]) for item in trades_by_day[date_value]) != int(
            row["pnl"]
        ):
            errors.append(f"{date_value}: trade cash flow")
        if sum(int(item["pnl"]) for item in campaigns_by_day[date_value]) != int(
            row["pnl"]
        ):
            errors.append(f"{date_value}: campaign pnl")

        quantities: dict[int, int] = defaultdict(int)
        for trade in trades_by_day[date_value]:
            campaign_id = int(trade["campaign"])
            quantity = int(trade["quantity"])
            if trade["side"] == "BUY":
                quantities[campaign_id] += quantity
                if quantity > math.floor(
                    int(trade["bar_volume"]) * VOLUME_PARTICIPATION
                ):
                    errors.append(f"{date_value}: entry participation")
            else:
                quantities[campaign_id] -= quantity
        if any(quantity != 0 for quantity in quantities.values()):
            errors.append(f"{date_value}: open quantity")
        previous_cash = int(row["end_cash"])

    return {
        "passed": not errors,
        "error_count": len(errors),
        "error_samples": errors[:20],
        "daily_rows": len(daily),
        "campaign_rows": len(campaigns),
        "trade_rows": len(trades),
    }


def run() -> dict[str, object]:
    profile = json.loads(base.PROFILE_PATH.read_text(encoding="utf-8"))
    groups = base.derive_expiry_groups(profile)
    files = base.target_file_map()
    volatility_dates, volatility_values = volatility_closes()
    cash = INITIAL_CASH
    daily_rows: list[dict[str, object]] = []
    campaign_rows: list[dict[str, object]] = []
    trades: list[dict[str, object]] = []

    for group in groups:
        date_value = group["expiry_date"]
        previous_vix = previous_volatility(
            date_value, volatility_dates, volatility_values
        )
        if previous_vix is None:
            raise RuntimeError(f"{date_value}: 직전 변동성지수 없음")
        contracts = base.load_expiry_day(files[date_value], group)
        day, day_campaigns, day_trades = simulate_day(
            date_value, contracts, previous_vix, cash
        )
        cash = day.end_cash
        daily_rows.append(asdict(day))
        campaign_rows.extend(day_campaigns)
        trades.extend(day_trades)

    profits = [int(row["pnl"]) for row in daily_rows]
    sorted_profits = sorted(profits, reverse=True)
    equity = [INITIAL_CASH] + [int(row["end_cash"]) for row in daily_rows]
    summary = {
        "strategy": STRATEGY_VERSION,
        "expiry_days": len(daily_rows),
        "high_volatility_days": sum(
            bool(row["high_volatility"]) for row in daily_rows
        ),
        "signal_days": sum(bool(row["signal_found"]) for row in daily_rows),
        "trade_days": sum(bool(row["trade_day"]) for row in daily_rows),
        "profitable_days": sum(value > 0 for value in profits),
        "losing_days": sum(value < 0 for value in profits),
        "no_trade_days": sum(value == 0 for value in profits),
        "initial_cash": INITIAL_CASH,
        "ending_cash": cash,
        "total_pnl": sum(profits),
        "maximum_day_pnl": max(profits),
        "minimum_day_pnl": min(profits),
        "max_drawdown": max_drawdown(equity),
        "pnl_excluding_best_1_day": sum(profits) - sum(sorted_profits[:1]),
        "pnl_excluding_best_3_days": sum(profits) - sum(sorted_profits[:3]),
        "pnl_excluding_best_5_days": sum(profits) - sum(sorted_profits[:5]),
        "total_campaigns": len(campaign_rows),
        "total_buy_principal": sum(
            int(row["buy_principal"]) for row in daily_rows
        ),
        "total_fees": sum(int(row["fees"]) for row in daily_rows),
        "expired_contracts": sum(
            int(row["expired_quantity"]) for row in daily_rows
        ),
        "expiry_residual_days": sum(
            int(row["expired_quantity"]) > 0 for row in daily_rows
        ),
    }

    analysis_campaigns = []
    for row in campaign_rows:
        enriched = dict(row)
        enriched["entry_time_bucket"] = entry_time_bucket(
            int(row["first_entry_minute"])
        )
        matching_day = next(
            item for item in daily_rows if item["date"] == row["date"]
        )
        enriched["vix_bucket"] = vix_bucket(
            float(matching_day["previous_vix"])
        )
        enriched["month"] = str(row["date"])[:6]
        analysis_campaigns.append(enriched)

    losing_days = [
        row for row in daily_rows if int(row["pnl"]) < 0
    ]
    loss_analysis = {
        "strategy": STRATEGY_VERSION,
        "losing_day_count": len(losing_days),
        "losing_campaign_count": sum(
            int(row["pnl"]) < 0 for row in campaign_rows
        ),
        "total_losing_day_loss": sum(
            int(row["pnl"]) for row in losing_days
        ),
        "worst_days": sorted(
            losing_days, key=lambda row: int(row["pnl"])
        )[:15],
        "by_exit_reason": grouped_losses(
            analysis_campaigns, "exit_reason"
        ),
        "by_call_put": grouped_losses(analysis_campaigns, "call_put"),
        "by_entry_time": grouped_losses(
            analysis_campaigns, "entry_time_bucket"
        ),
        "by_vix": grouped_losses(analysis_campaigns, "vix_bucket"),
        "by_tranche_count": grouped_losses(
            analysis_campaigns, "tranche_count"
        ),
        "by_month": grouped_losses(analysis_campaigns, "month"),
    }
    audit_result = audit(daily_rows, campaign_rows, trades)
    result = {
        "status": "experimental_fixed_strategy",
        "strategy_version": STRATEGY_VERSION,
        "summary": summary,
        "loss_analysis": loss_analysis,
        "audit": audit_result,
        "limitations": [
            "전체 보유 자료에서 최초 실험한 전략이며 미사용 홀드아웃이 아닙니다.",
            "15:19까지 미체결된 잔량은 0원으로 평가했습니다.",
            "분봉 거래량 참여율은 실제 호가 체결을 보장하지 않습니다.",
        ],
    }

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    write_csv(OUTPUT_DIR / "strategy1_daily.csv", daily_rows)
    write_csv(OUTPUT_DIR / "strategy1_campaigns.csv", campaign_rows)
    write_csv(OUTPUT_DIR / "strategy1_trades.csv", trades)
    (OUTPUT_DIR / "strategy1_summary.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (OUTPUT_DIR / "strategy1_loss_analysis.json").write_text(
        json.dumps(loss_analysis, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (OUTPUT_DIR / "strategy1_audit.json").write_text(
        json.dumps(audit_result, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if not audit_result["passed"]:
        raise SystemExit(1)
    return result


if __name__ == "__main__":
    run()
