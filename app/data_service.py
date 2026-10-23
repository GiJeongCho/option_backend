from __future__ import annotations

import csv
import math
import re
from datetime import date, datetime, time, timedelta, timezone
from functools import lru_cache
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook

from .index_minute_catalog import IndexMinuteCatalog


Session = Literal["day", "night"]

DATE_PATTERN = re.compile(r"^\d{8}$")
FILE_DATE_PATTERN = re.compile(r"_(\d{8})\.csv$")
SERIES_PATTERN = re.compile(r"\b(\d{4}W\d)\b")
KST = timezone(timedelta(hours=9))
SESSION_MARKET_ID: dict[Session, str] = {"day": "DRV", "night": "NDV"}
OPTION_DATASET_MARKER = "[파생 옵션 코스피200 위클리]일중 매매정보(1분)"


class DataNotFoundError(LookupError):
    pass


def _to_float(value: str | None) -> float | None:
    if value is None or value.strip() == "":
        return None
    try:
        return float(value)
    except ValueError:
        return None


def _to_int(value: str | None) -> int:
    number = _to_float(value)
    return int(number) if number is not None else 0


def _series(name: str) -> str:
    match = SERIES_PATTERN.search(name)
    return match.group(1) if match else ""


def _product_kind(name: str) -> str:
    return "MONDAY" if "위클리M" in name else "THURSDAY"


def _session_sort_key(hhmm: str, session: Session) -> int:
    value = int(hhmm.strip().zfill(4))
    if session == "night" and value >= 1800:
        return value - 2400
    return value


class MarketDataService:
    def __init__(self, backend_dir: Path | None = None) -> None:
        self.backend_dir = backend_dir or Path(__file__).resolve().parents[1]
        self.rowdata_dir = self.backend_dir / "rowdata"
        self.volatility_path = self.rowdata_dir / "코스피200 변동성지수.xlsx"
        self.index_minute_catalog = IndexMinuteCatalog(self.rowdata_dir)
        self.spot_minute_path = self.index_minute_catalog.cache_path("spot")
        self.futures_minute_path = self.index_minute_catalog.cache_path(
            "futures"
        )

    def option_file_index(self) -> dict[str, Path]:
        index: dict[str, Path] = {}
        if not self.rowdata_dir.exists():
            return index

        for directory in self.rowdata_dir.iterdir():
            if (
                not directory.is_dir()
                or OPTION_DATASET_MARKER not in directory.name
            ):
                continue
            for path in directory.glob("*.csv"):
                match = FILE_DATE_PATTERN.search(path.name)
                if match is None:
                    continue
                trading_date = match.group(1)
                if trading_date in index:
                    raise RuntimeError(f"중복 옵션 파일 날짜: {trading_date}")
                index[trading_date] = path
        return index

    def list_option_dates(self) -> list[str]:
        return sorted(self.option_file_index())

    def _option_path(self, trading_date: str) -> Path:
        if DATE_PATTERN.fullmatch(trading_date) is None:
            raise ValueError("날짜는 YYYYMMDD 형식이어야 합니다.")
        try:
            return self.option_file_index()[trading_date]
        except KeyError as error:
            raise DataNotFoundError(
                f"{trading_date} 옵션 파일을 찾을 수 없습니다."
            ) from error

    @lru_cache(maxsize=32)
    def list_contracts(
        self, trading_date: str, session: Session = "day"
    ) -> list[dict[str, object]]:
        path = self._option_path(trading_date)
        market_id = SESSION_MARKET_ID[session]
        contracts: dict[str, dict[str, object]] = {}

        with path.open("r", encoding="cp949", newline="") as source:
            for row in csv.DictReader(source):
                if row["시장ID"] != market_id:
                    continue

                open_price = _to_float(row["시가"])
                high = _to_float(row["고가"])
                low = _to_float(row["저가"])
                close = _to_float(row["종가"])
                if None in (open_price, high, low, close):
                    continue

                code = row["종목코드"]
                normalized_time = row["기준시각"].strip().zfill(4)
                sort_key = _session_sort_key(normalized_time, session)
                current = contracts.get(code)
                if current is None:
                    current = {
                        "code": code,
                        "name": row["종목명"],
                        "series": _series(row["종목명"]),
                        "productKind": _product_kind(row["종목명"]),
                        "callPut": row["콜풋구분"],
                        "strike": _to_float(row["행사가격"]),
                        "barCount": 0,
                        "totalVolume": 0,
                        "minLow": low,
                        "maxHigh": high,
                        "firstOpen": open_price,
                        "lastClose": close,
                        "firstTime": normalized_time,
                        "lastTime": normalized_time,
                        "touchedPremium1To5": False,
                        "_firstSortKey": sort_key,
                        "_lastSortKey": sort_key,
                    }
                    contracts[code] = current

                current["barCount"] = int(current["barCount"]) + 1
                current["totalVolume"] = int(current["totalVolume"]) + _to_int(
                    row["거래량"]
                )
                current["minLow"] = min(float(current["minLow"]), low)
                current["maxHigh"] = max(float(current["maxHigh"]), high)

                if sort_key < int(current["_firstSortKey"]):
                    current["_firstSortKey"] = sort_key
                    current["firstTime"] = normalized_time
                    current["firstOpen"] = open_price
                if sort_key >= int(current["_lastSortKey"]):
                    current["_lastSortKey"] = sort_key
                    current["lastTime"] = normalized_time
                    current["lastClose"] = close
                if 1.0 <= low <= 5.0:
                    current["touchedPremium1To5"] = True

        public_contracts = [
            {
                key: value
                for key, value in contract.items()
                if not key.startswith("_")
            }
            for contract in contracts.values()
        ]
        return sorted(
            public_contracts,
            key=lambda item: (
                str(item["productKind"]),
                str(item["callPut"]),
                float(item["strike"] or 0),
            ),
        )

    @staticmethod
    def _bar_timestamp(
        trading_date: str, hhmm: str, session: Session
    ) -> tuple[int, str, int]:
        trading_day = datetime.strptime(trading_date, "%Y%m%d").date()
        normalized = hhmm.strip().zfill(4)
        hour = int(normalized[:2])
        minute = int(normalized[2:])

        calendar_day = trading_day
        sort_key = hour * 100 + minute
        if session == "night" and hour >= 18:
            calendar_day -= timedelta(days=1)
            sort_key -= 2400

        moment = datetime.combine(calendar_day, time(hour, minute), tzinfo=KST)
        return int(moment.timestamp()), moment.strftime("%m-%d %H:%M"), sort_key

    @lru_cache(maxsize=128)
    def option_bars(
        self, trading_date: str, code: str, session: Session = "day"
    ) -> list[dict[str, object]]:
        path = self._option_path(trading_date)
        market_id = SESSION_MARKET_ID[session]
        bars: dict[int, dict[str, object]] = {}

        with path.open("r", encoding="cp949", newline="") as source:
            for row in csv.DictReader(source):
                if row["시장ID"] != market_id or row["종목코드"] != code:
                    continue

                open_price = _to_float(row["시가"])
                high = _to_float(row["고가"])
                low = _to_float(row["저가"])
                close = _to_float(row["종가"])
                if None in (open_price, high, low, close):
                    continue

                timestamp, label, sort_key = self._bar_timestamp(
                    trading_date, row["기준시각"], session
                )
                current = bars.get(timestamp)
                if current is None:
                    bars[timestamp] = {
                        "time": timestamp,
                        "label": label,
                        "sortKey": sort_key,
                        "open": open_price,
                        "high": high,
                        "low": low,
                        "close": close,
                        "volume": _to_int(row["거래량"]),
                    }
                    continue

                current["high"] = max(float(current["high"]), high)
                current["low"] = min(float(current["low"]), low)
                current["close"] = close
                current["volume"] = int(current["volume"]) + _to_int(row["거래량"])

        if not bars:
            raise DataNotFoundError(
                f"{trading_date} {code} {session} 분봉을 찾을 수 없습니다."
            )
        return sorted(bars.values(), key=lambda item: int(item["time"]))

    @lru_cache(maxsize=128)
    def option_bar_grid(
        self, trading_date: str, code: str
    ) -> list[dict[str, object]]:
        actual = self.option_bars(trading_date, code, "day")
        by_minute = {int(row["sortKey"]): row for row in actual}
        result: list[dict[str, object]] = []
        previous_close: float | None = None
        for total in range(8 * 60 + 45, 15 * 60 + 20):
            hour, minute = divmod(total, 60)
            hhmm = hour * 100 + minute
            row = by_minute.get(hhmm)
            if row is not None:
                current = dict(row)
                current["synthetic"] = False
                result.append(current)
                previous_close = float(row["close"])
                continue
            if previous_close is None:
                continue
            timestamp, label, sort_key = self._bar_timestamp(
                trading_date, f"{hhmm:04d}", "day"
            )
            result.append(
                {
                    "time": timestamp,
                    "label": label,
                    "sortKey": sort_key,
                    "open": previous_close,
                    "high": previous_close,
                    "low": previous_close,
                    "close": previous_close,
                    "volume": 0,
                    "synthetic": True,
                }
            )
        return result

    @lru_cache(maxsize=1)
    def volatility_bars(self) -> list[dict[str, object]]:
        if not self.volatility_path.exists():
            raise DataNotFoundError(
                f"변동성지수 파일을 찾을 수 없습니다: {self.volatility_path}"
            )

        workbook = load_workbook(
            self.volatility_path, read_only=True, data_only=True
        )
        try:
            sheet = workbook["코스피200 변동성지수"]
            result: list[dict[str, object]] = []
            for row in sheet.iter_rows(min_row=2, max_col=5, values_only=True):
                raw_date, raw_open, raw_high, raw_low, raw_close = row
                if raw_date is None or None in (
                    raw_open,
                    raw_high,
                    raw_low,
                    raw_close,
                ):
                    continue

                if isinstance(raw_date, datetime):
                    trading_day = raw_date.date()
                elif isinstance(raw_date, date):
                    trading_day = raw_date
                else:
                    text = str(raw_date).strip()
                    parsed = None
                    for date_format in ("%Y-%m-%d", "%Y%m%d", "%Y.%m.%d"):
                        try:
                            parsed = datetime.strptime(text, date_format).date()
                            break
                        except ValueError:
                            continue
                    if parsed is None:
                        continue
                    trading_day = parsed

                result.append(
                    {
                        "date": trading_day.isoformat(),
                        "open": float(raw_open),
                        "high": float(raw_high),
                        "low": float(raw_low),
                        "close": float(raw_close),
                    }
                )
            return sorted(result, key=lambda item: str(item["date"]))
        finally:
            workbook.close()

    @staticmethod
    def _normalized_slope(values: list[float]) -> float | None:
        if len(values) < 3 or abs(values[-1]) <= 1e-12:
            return None
        count = len(values)
        mean_x = (count - 1) / 2.0
        mean_y = sum(values) / count
        denominator = sum(
            (index - mean_x) ** 2 for index in range(count)
        )
        if denominator <= 1e-12:
            return None
        numerator = sum(
            (index - mean_x) * (value - mean_y)
            for index, value in enumerate(values)
        )
        return numerator / denominator / values[-1]

    @staticmethod
    def _realized_volatility(
        values: list[float],
        period: int,
    ) -> float | None:
        if len(values) < period + 1:
            return None
        selected = values[-(period + 1) :]
        if any(value <= 0 for value in selected):
            return None
        returns = [
            math.log(right / left)
            for left, right in zip(selected, selected[1:])
        ]
        if len(returns) < 2:
            return 0.0
        mean = sum(returns) / len(returns)
        variance = sum((value - mean) ** 2 for value in returns) / (
            len(returns) - 1
        )
        return math.sqrt(max(0.0, variance))

    @staticmethod
    def _market_minute_timestamp(
        raw_value: str,
    ) -> tuple[str, int, str, int] | None:
        try:
            moment = datetime.strptime(
                raw_value.strip(),
                "%Y-%m-%d %H:%M",
            ).replace(tzinfo=KST)
        except ValueError:
            return None
        minute = moment.hour * 100 + moment.minute
        return (
            moment.strftime("%Y%m%d"),
            int(moment.timestamp()),
            moment.strftime("%m-%d %H:%M"),
            minute,
        )

    @staticmethod
    @lru_cache(maxsize=8)
    def _market_date_offsets(
        path_text: str,
        file_size: int,
        modified_ns: int,
    ) -> dict[str, tuple[int, int]]:
        del file_size, modified_ns
        offsets: dict[str, tuple[int, int]] = {}
        current_date = ""
        current_start = 0
        path = Path(path_text)
        with path.open("rb") as source:
            while True:
                start = source.tell()
                line = source.readline()
                if not line:
                    if current_date:
                        offsets[current_date] = (current_start, start)
                    break
                prefix = line[:10].decode("ascii", errors="ignore")
                if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", prefix):
                    continue
                date_value = prefix.replace("-", "")
                if date_value == current_date:
                    continue
                if current_date:
                    offsets[current_date] = (current_start, start)
                current_date = date_value
                current_start = start
        return offsets

    @lru_cache(maxsize=64)
    def market_index_bars(
        self,
        trading_date: str,
        source: Literal["spot", "futures"],
    ) -> list[dict[str, object]]:
        if DATE_PATTERN.fullmatch(trading_date) is None:
            raise ValueError("날짜는 YYYYMMDD 형식이어야 합니다.")
        try:
            path = self.index_minute_catalog.normalized_path(source)
        except (FileNotFoundError, ValueError) as error:
            raise DataNotFoundError(str(error)) from error

        stat = path.stat()
        offsets = self._market_date_offsets(
            str(path),
            stat.st_size,
            stat.st_mtime_ns,
        )
        byte_range = offsets.get(trading_date)
        if byte_range is None:
            raise DataNotFoundError(
                f"{trading_date} {source} 지수 분봉을 찾을 수 없습니다."
            )
        with path.open("rb") as raw:
            raw.seek(byte_range[0])
            payload = raw.read(byte_range[1] - byte_range[0]).decode(
                "utf-8-sig"
            )

        bars: dict[int, dict[str, object]] = {}
        for values in csv.reader(payload.splitlines()):
            if len(values) < 5:
                continue
            parsed = self._market_minute_timestamp(values[0])
            if parsed is None:
                continue
            _, timestamp, label, minute = parsed
            if not 845 <= minute <= 1520:
                continue
            prices = [
                _to_float(str(value).replace(",", ""))
                for value in values[1:5]
            ]
            if any(value is None for value in prices):
                continue
            open_price, high, low, close = prices
            bars[timestamp] = {
                "time": timestamp,
                "label": label,
                "sortKey": minute,
                "open": open_price,
                "high": high,
                "low": low,
                "close": close,
            }
        if not bars:
            raise DataNotFoundError(
                f"{trading_date} {source} 지수 분봉을 찾을 수 없습니다."
            )
        result = sorted(bars.values(), key=lambda item: int(item["time"]))
        baseline = float(result[0]["close"])
        for row in result:
            row["normalizedClose"] = (
                float(row["close"]) / baseline * 100.0
                if baseline > 0
                else 100.0
            )
        return result

    def market_context(
        self,
        trading_date: str,
        decision_minute: int,
    ) -> dict[str, object]:
        if not 845 <= decision_minute <= 1520:
            raise ValueError("판단시각은 08:45~15:20이어야 합니다.")
        spot = self.market_index_bars(trading_date, "spot")
        futures = self.market_index_bars(trading_date, "futures")
        spot_by_time = {int(row["time"]): row for row in spot}
        futures_by_time = {int(row["time"]): row for row in futures}
        common_times = sorted(spot_by_time.keys() & futures_by_time.keys())
        if common_times:
            first_common = common_times[0]
            spot_baseline = float(spot_by_time[first_common]["close"])
            futures_baseline = float(futures_by_time[first_common]["close"])
            for row in spot:
                row["normalizedClose"] = (
                    float(row["close"]) / spot_baseline * 100.0
                )
            for row in futures:
                row["normalizedClose"] = (
                    float(row["close"]) / futures_baseline * 100.0
                )
        basis = [
            {
                "time": timestamp,
                "label": spot_by_time[timestamp]["label"],
                "value": (
                    float(futures_by_time[timestamp]["close"])
                    - float(spot_by_time[timestamp]["close"])
                ),
            }
            for timestamp in common_times
        ]

        selected_iso = datetime.strptime(
            trading_date,
            "%Y%m%d",
        ).date().isoformat()
        vix_source = [
            dict(row)
            for row in self.volatility_bars()
            if str(row["date"]) < selected_iso
        ][-120:]
        closes: list[float] = []
        vix_bars: list[dict[str, object]] = []
        for row in vix_source:
            closes.append(float(row["close"]))
            current = dict(row)
            current["ma5"] = (
                sum(closes[-5:]) / 5 if len(closes) >= 5 else None
            )
            current["ma20"] = (
                sum(closes[-20:]) / 20 if len(closes) >= 20 else None
            )
            vix_bars.append(current)

        known_spot = [
            float(row["close"])
            for row in spot
            if int(row["sortKey"]) <= decision_minute
        ]
        known_futures = [
            float(row["close"])
            for row in futures
            if int(row["sortKey"]) <= decision_minute
        ]
        latest_vix = vix_bars[-1] if vix_bars else None
        previous_vix_close = closes[-2] if len(closes) >= 2 else None
        vix_ma5 = (
            sum(closes[-5:]) / 5 if len(closes) >= 5 else None
        )
        vix_ma20 = (
            sum(closes[-20:]) / 20 if len(closes) >= 20 else None
        )
        vix_slope_5 = self._normalized_slope(closes[-5:])
        vix_slope_20 = self._normalized_slope(closes[-20:])
        vix_slope_120 = self._normalized_slope(closes[-120:])

        daily_points = 0
        if latest_vix is not None:
            latest_close = float(latest_vix["close"])
            daily_points += int(vix_ma5 is not None and latest_close > vix_ma5)
            daily_points += int(vix_slope_5 is not None and vix_slope_5 > 0)
            daily_points += int(
                previous_vix_close is not None
                and latest_close > previous_vix_close
            )
        spot_rv = {
            period: self._realized_volatility(known_spot, period)
            for period in (5, 20, 60)
        }
        futures_rv = {
            period: self._realized_volatility(known_futures, period)
            for period in (5, 20, 60)
        }
        intraday_points = sum(
            (
                int(
                    spot_rv[5] is not None
                    and spot_rv[20] is not None
                    and spot_rv[5] > spot_rv[20]
                ),
                int(
                    futures_rv[5] is not None
                    and futures_rv[20] is not None
                    and futures_rv[5] > futures_rv[20]
                ),
                int(
                    spot_rv[20] is not None
                    and spot_rv[60] is not None
                    and spot_rv[20] > spot_rv[60]
                ),
                int(
                    futures_rv[20] is not None
                    and futures_rv[60] is not None
                    and futures_rv[20] > futures_rv[60]
                ),
            )
        )
        if daily_points >= 2 and intraday_points >= 1:
            regime = "EXPANDING"
        elif (
            vix_slope_5 is not None
            and vix_slope_5 < 0
            and latest_vix is not None
            and previous_vix_close is not None
            and float(latest_vix["close"]) < previous_vix_close
            and intraday_points == 0
        ):
            regime = "CONTRACTING"
        elif (
            latest_vix is not None
            and vix_ma20 is not None
            and float(latest_vix["close"]) > vix_ma20
            and vix_slope_20 is not None
            and vix_slope_20 > 0
        ):
            regime = "HIGH_STABLE"
        else:
            regime = "NORMAL"

        spot_metadata = self.index_minute_catalog.metadata("spot")
        futures_metadata = self.index_minute_catalog.metadata("futures")
        return {
            "date": trading_date,
            "decisionMinute": decision_minute,
            "spot": spot,
            "futures": futures,
            "basis": basis,
            "vix": {
                "knownThrough": (
                    latest_vix["date"] if latest_vix is not None else None
                ),
                "bars": vix_bars,
                "snapshot": {
                    "close": (
                        latest_vix["close"]
                        if latest_vix is not None
                        else None
                    ),
                    "ma5": vix_ma5,
                    "ma20": vix_ma20,
                    "slope5": vix_slope_5,
                    "slope20": vix_slope_20,
                    "slope120": vix_slope_120,
                },
            },
            "regime": {
                "key": regime,
                "vixTrendPoints": daily_points,
                "intradayExpansionPoints": intraday_points,
                "score": daily_points + intraday_points,
                "spotRv5": spot_rv[5],
                "spotRv20": spot_rv[20],
                "spotRv60": spot_rv[60],
                "futuresRv5": futures_rv[5],
                "futuresRv20": futures_rv[20],
                "futuresRv60": futures_rv[60],
            },
            "indexData": {
                "precedence": (
                    "신규 통합 XLSX 우선, 기존 CSV는 누락 시각과 "
                    "2026-09-30 이후 구간 보충"
                ),
                "spot": {
                    "start": spot_metadata["start"],
                    "end": spot_metadata["end"],
                    "rows": spot_metadata["mergedRows"],
                },
                "futures": {
                    "start": futures_metadata["start"],
                    "end": futures_metadata["end"],
                    "rows": futures_metadata["mergedRows"],
                },
            },
            "dataNotice": (
                "VIX는 일봉이므로 거래일보다 앞선 마지막 확정값만 "
                "표시합니다. 현물·선물은 통합 XLSX를 우선하고 기존 "
                "CSV로 누락 시각과 이후 구간을 보충한 1분봉입니다."
            ),
        }
