from __future__ import annotations

import csv
import json
import math
import os
from datetime import date, datetime
from pathlib import Path
from typing import Iterator, Literal

from openpyxl import load_workbook


IndexSource = Literal["spot", "futures"]

CACHE_SCHEMA_VERSION = 1
SOURCE_FILES: dict[IndexSource, tuple[str, str]] = {
    "spot": (
        "코스피200_현물지수_1분_20230801_20260930.xlsx",
        "코스피 200현물지수_1분봉.csv",
    ),
    "futures": (
        "코스피200_선물지수_1분_20230801_20260930.xlsx",
        "코스피 200선물지수_1분봉.csv",
    ),
}


def _timestamp_text(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time()).strftime(
            "%Y-%m-%d %H:%M"
        )
    text = str(value or "").strip()
    if not text:
        return None
    for pattern in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(text, pattern).strftime("%Y-%m-%d %H:%M")
        except ValueError:
            continue
    return None


def _number(value: object) -> float | None:
    if value is None:
        return None
    try:
        number = float(str(value).replace(",", ""))
    except ValueError:
        return None
    return number if math.isfinite(number) else None


def _validated_ohlc(
    values: tuple[object, object, object, object],
    source_name: str,
    timestamp: str,
) -> tuple[float, float, float, float]:
    parsed = tuple(_number(value) for value in values)
    if any(value is None for value in parsed):
        raise ValueError(f"{source_name} {timestamp}: OHLC 값 누락")
    open_price, high, low, close = parsed
    assert open_price is not None
    assert high is not None
    assert low is not None
    assert close is not None
    if (
        min(open_price, high, low, close) <= 0
        or high + 1e-9 < max(open_price, close)
        or low - 1e-9 > min(open_price, close)
        or high + 1e-9 < low
    ):
        raise ValueError(f"{source_name} {timestamp}: OHLC 범위 오류")
    return open_price, high, low, close


class IndexMinuteCatalog:
    def __init__(self, rowdata_dir: Path) -> None:
        self.rowdata_dir = rowdata_dir
        self.cache_dir = rowdata_dir / ".cache" / "market_index"

    def cache_path(self, source: IndexSource) -> Path:
        return self.cache_dir / f"{source}_minute_merged.csv"

    def manifest_path(self, source: IndexSource) -> Path:
        return self.cache_dir / f"{source}_minute_merged.json"

    def input_paths(self, source: IndexSource) -> list[Path]:
        return [
            self.rowdata_dir / name
            for name in SOURCE_FILES[source]
            if (self.rowdata_dir / name).exists()
        ]

    def _input_signature(self, source: IndexSource) -> list[dict[str, object]]:
        return [
            {
                "name": path.name,
                "size": path.stat().st_size,
                "modifiedNs": path.stat().st_mtime_ns,
            }
            for path in self.input_paths(source)
        ]

    def _cache_is_current(self, source: IndexSource) -> bool:
        cache_path = self.cache_path(source)
        manifest_path = self.manifest_path(source)
        if not cache_path.exists() or not manifest_path.exists():
            return False
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return False
        return (
            manifest.get("schemaVersion") == CACHE_SCHEMA_VERSION
            and manifest.get("inputs") == self._input_signature(source)
        )

    @staticmethod
    def _workbook_rows(
        path: Path,
    ) -> Iterator[tuple[str, tuple[float, float, float, float]]]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            for row_number, row in enumerate(
                sheet.iter_rows(values_only=True),
                start=1,
            ):
                if row_number == 1:
                    continue
                timestamp = _timestamp_text(row[0] if row else None)
                if timestamp is None:
                    continue
                if len(row) < 5:
                    raise ValueError(
                        f"{path.name} {timestamp}: 예상하지 못한 열 수"
                    )
                yield timestamp, _validated_ohlc(
                    (row[1], row[2], row[3], row[4]),
                    path.name,
                    timestamp,
                )
        finally:
            workbook.close()

    @staticmethod
    def _legacy_csv_rows(
        path: Path,
    ) -> Iterator[tuple[str, tuple[float, float, float, float]]]:
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            for row in csv.reader(source):
                if len(row) < 5:
                    continue
                timestamp = _timestamp_text(row[0])
                if timestamp is None:
                    continue
                yield timestamp, _validated_ohlc(
                    (row[1], row[2], row[3], row[4]),
                    path.name,
                    timestamp,
                )

    def _build(self, source: IndexSource) -> None:
        xlsx_name, csv_name = SOURCE_FILES[source]
        xlsx_path = self.rowdata_dir / xlsx_name
        csv_path = self.rowdata_dir / csv_name
        if not xlsx_path.exists() and not csv_path.exists():
            raise FileNotFoundError(
                f"{source} 현물·선물 분봉 원자료를 찾을 수 없습니다."
            )

        records: dict[str, tuple[float, float, float, float]] = {}
        workbook_rows = 0
        workbook_duplicates = 0
        if xlsx_path.exists():
            for timestamp, ohlc in self._workbook_rows(xlsx_path):
                workbook_rows += 1
                workbook_duplicates += int(timestamp in records)
                records[timestamp] = ohlc

        csv_rows = 0
        overlap_rows = 0
        overlap_exact_rows = 0
        overlap_close_abs_max = 0.0
        csv_filled_rows = 0
        if csv_path.exists():
            for timestamp, ohlc in self._legacy_csv_rows(csv_path):
                csv_rows += 1
                workbook_value = records.get(timestamp)
                if workbook_value is not None:
                    overlap_rows += 1
                    overlap_exact_rows += int(
                        all(
                            abs(left - right) <= 1e-9
                            for left, right in zip(workbook_value, ohlc)
                        )
                    )
                    overlap_close_abs_max = max(
                        overlap_close_abs_max,
                        abs(workbook_value[3] - ohlc[3]),
                    )
                    continue
                records[timestamp] = ohlc
                csv_filled_rows += 1

        if not records:
            raise ValueError(f"{source}: 유효한 현물·선물 분봉이 없습니다.")

        ordered_timestamps = sorted(records)
        self.cache_dir.mkdir(parents=True, exist_ok=True)
        cache_path = self.cache_path(source)
        manifest_path = self.manifest_path(source)
        cache_temp = cache_path.with_name(
            f".{cache_path.name}.{os.getpid()}.tmp"
        )
        manifest_temp = manifest_path.with_name(
            f".{manifest_path.name}.{os.getpid()}.tmp"
        )
        with cache_temp.open("w", encoding="utf-8-sig", newline="") as target:
            writer = csv.writer(target, lineterminator="\n")
            writer.writerow(["통합 시장지수 데이터", "", "", "", ""])
            writer.writerow(["날짜", "시가", "고가", "저가", "종가"])
            for timestamp in ordered_timestamps:
                writer.writerow((timestamp, *records[timestamp]))

        manifest = {
            "schemaVersion": CACHE_SCHEMA_VERSION,
            "source": source,
            "precedence": "XLSX 우선, 기존 CSV는 XLSX 누락 시각과 이후 구간 보충",
            "inputs": self._input_signature(source),
            "workbookRows": workbook_rows,
            "workbookDuplicateTimestamps": workbook_duplicates,
            "legacyCsvRows": csv_rows,
            "overlapRows": overlap_rows,
            "overlapExactOhlcRows": overlap_exact_rows,
            "overlapExactOhlcPct": (
                round(overlap_exact_rows / overlap_rows * 100.0, 6)
                if overlap_rows
                else None
            ),
            "overlapCloseAbsoluteMax": overlap_close_abs_max,
            "legacyCsvFilledRows": csv_filled_rows,
            "mergedRows": len(records),
            "start": ordered_timestamps[0],
            "end": ordered_timestamps[-1],
        }
        manifest_temp.write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        cache_temp.replace(cache_path)
        manifest_temp.replace(manifest_path)

    def normalized_path(self, source: IndexSource) -> Path:
        if not self._cache_is_current(source):
            self._build(source)
        return self.cache_path(source)

    def metadata(self, source: IndexSource) -> dict[str, object]:
        self.normalized_path(source)
        return json.loads(
            self.manifest_path(source).read_text(encoding="utf-8")
        )
